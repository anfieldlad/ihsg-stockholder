import os
import openpyxl
import pytest

MODEL_PATH = '/home/hermes/company/ihsg/freemium-model-v3.1.xlsx'
PLAN_PATH = '/home/hermes/company/ihsg/freemium-plan-v3.1.md'
REPO_MODEL_PATH = 'docs/freemium-model-v3.1.xlsx'
REPO_PLAN_PATH = 'docs/freemium-plan-v3.1.md'

EXPECTED_SHEETS = [
    'Executive_Summary',
    'Model_Parameters',
    'Skenario_Cold_Start',
    'Skenario_Konservatif',
    'Skenario_Moderat_Base',
    'Skenario_Optimis',
    'OneTime_vs_Recurring',
    'Comp_Accounts_ROI',
    'Perbandingan_Skenario'
]

def test_file_existence():
    """Verify deliverables exist in company directory and in repo docs/."""
    assert os.path.exists(MODEL_PATH), f"Missing {MODEL_PATH}"
    assert os.path.exists(PLAN_PATH), f"Missing {PLAN_PATH}"
    assert os.path.exists(REPO_MODEL_PATH), f"Missing {REPO_MODEL_PATH}"
    assert os.path.exists(REPO_PLAN_PATH), f"Missing {REPO_PLAN_PATH}"

def test_workbook_structure():
    """Verify all 9 sheets are present and fullCalcOnLoad is enabled."""
    wb = openpyxl.load_workbook(MODEL_PATH, data_only=False)
    assert wb.sheetnames == EXPECTED_SHEETS
    assert wb.calculation.fullCalcOnLoad is True

def test_prices_untouched():
    """Verify Bobby's prices remain unchanged in Model_Parameters."""
    wb = openpyxl.load_workbook(MODEL_PATH, data_only=False)
    ws = wb['Model_Parameters']
    assert ws['B7'].value == 19000      # Investor 1M
    assert ws['B8'].value == 190000     # Investor 12M
    assert ws['B9'].value == 49000      # Pakar 1M
    assert ws['B10'].value == 490000    # Pakar 12M
    assert ws['B11'].value == 599000    # Founder Pass Lifetime
    assert ws['B23'].value == 5         # Comp accounts cap

def test_supabase_line_dropped_and_fixed_costs():
    """Verify Supabase Pro line is dropped and Phase 0 fixed cost is Rp 390,000."""
    wb = openpyxl.load_workbook(MODEL_PATH, data_only=False)
    ws = wb['Model_Parameters']
    # Row 42: Database line should be 0 across all tiers
    assert ws['B42'].value == 0
    assert ws['C42'].value == 0
    assert ws['D42'].value == 0
    assert ws['E42'].value == 0
    assert "SQLite" in str(ws['A42'].value)
    
    # Phase 0 tech infra = VPS 70k (20k domain + 50k server)
    assert ws['B37'].value == 20000
    assert ws['B38'].value == 50000
    # AI Claude = 320k
    assert ws['B40'].value == 320000

def test_comp_accounts_excluded_from_paying_trigger():
    """Verify comp accounts are strictly excluded from paying users trigger."""
    wb = openpyxl.load_workbook(MODEL_PATH, data_only=False)
    ws_comp = wb['Comp_Accounts_ROI']
    sql = str(ws_comp['C24'].value)
    assert "source = 'paid'" in sql or "source='paid'" in sql
    assert "amount_idr > 0" in sql or "amount_idr>0" in sql
    assert "COUNT(DISTINCT" in sql

    # Verify scenario sheets use paying users (row 31), not total users (row 33)
    for sname in ['Skenario_Cold_Start', 'Skenario_Konservatif', 'Skenario_Moderat_Base', 'Skenario_Optimis']:
        ws = wb[sname]
        # Row 34 trigger formula should reference row 31
        assert ws['B34'].value == '=B31'
        assert ws['C34'].value == '=C31'

def test_plan_v3_1_content():
    """Verify freemium-plan-v3.1.md confirms Jim D8 rules and specifies the report."""
    with open(PLAN_PATH, 'r') as f:
        content = f.read()
    
    assert "Same Tier Stacks" in content or "R1" in content
    assert "Lower Tier Blocked" in content or "R2" in content
    assert "GRACE_DAYS = 0" in content or "R7" in content
    assert "Upgrade" in content or "R4" in content
    assert "Founder Pass" in content or "R5" in content
    assert "python -m ihsg.admin report" in content
    assert "COUNT(DISTINCT e.user_id)" in content
